"""
Батч-режим: прогнать список товаров из Excel через TenderSearchSystem и
записать результаты (рыночная цена / статус / уверенность) обратно в копию файла
с цветовой заливкой по статусу.

Стиль оформления перенесён из tender_claude_code/excel_handler.py. Номера колонок
входного файла НЕ хардкодятся (в отличие от excel_handler.py) — эталонного файла
для батч-режима пока нет, поэтому они передаются явно через CLI-флаги.
"""

from __future__ import annotations

import logging
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import EXCEL_FILLS, EXCEL_HEADER_FILL, EXCEL_HEADER_FONT_COLOR, NEW_EXCEL_HEADERS
from .orchestrator import TenderSearchSystem

log = logging.getLogger("tender")

_HEADER_FILL = PatternFill("solid", fgColor=EXCEL_HEADER_FILL)
_HEADER_FONT = Font(color=EXCEL_HEADER_FONT_COLOR, bold=True, size=10)
_BOLD = Font(bold=True)
_THIN = Side(style="thin", color="CCCCCC")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_COL_WIDTHS = [16, 14, 14, 24, 14, 14]


def _cell_str(ws, row: int, col: int) -> str:
    v = ws.cell(row=row, column=col).value
    return str(v).strip() if v is not None else ""


def _cell_float(ws, row: int, col: int) -> Optional[float]:
    v = ws.cell(row=row, column=col).value
    try:
        return float(v) if v is not None else None
    except (TypeError, ValueError):
        return None


def _fmt_price(v: Optional[float]) -> str:
    if v is None:
        return "—"
    return f"{v:,.0f}".replace(",", " ")


def read_tender_excel(path: str, name_col: int, unit_col: Optional[int],
                       actual_price_col: Optional[int],
                       limit: Optional[int] = None) -> tuple[list[dict], "openpyxl.Workbook"]:
    """Возвращает (rows, workbook). rows: [{excel_row, name, unit, actual_price}, ...]."""
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    rows: list[dict] = []
    for row_num in range(2, ws.max_row + 1):  # строка 1 — заголовок
        name = _cell_str(ws, row_num, name_col)
        if not name:
            continue
        rows.append({
            "excel_row": row_num,
            "name": name,
            "unit": _cell_str(ws, row_num, unit_col) if unit_col else "",
            "actual_price": _cell_float(ws, row_num, actual_price_col) if actual_price_col else None,
        })
        if limit and len(rows) >= limit:
            break

    log.info("Батч-режим: прочитано %d строк из %s", len(rows), path)
    return rows, wb


def run_batch(system: TenderSearchSystem, input_path: str, output_path: str,
              name_col: int, unit_col: Optional[int] = None,
              actual_price_col: Optional[int] = None,
              limit: Optional[int] = None) -> None:
    """Прогоняет все строки входного файла через query_structured() и
    сохраняет результат в output_path с цветовой заливкой по статусу."""
    rows, wb = read_tender_excel(input_path, name_col, unit_col, actual_price_col, limit)
    ws = wb.active

    first_new_col = ws.max_column + 1
    for offset, label in enumerate(NEW_EXCEL_HEADERS):
        col = first_new_col + offset
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col)].width = (
            _COL_WIDTHS[offset] if offset < len(_COL_WIDTHS) else 14
        )

    for i, row in enumerate(rows, 1):
        query_text = f"{row['name']} {row['unit']}".strip() if row["unit"] else row["name"]
        outcome = system.query_structured(query_text, want_status_against=row["actual_price"])
        print(f"  [{i}/{len(rows)}] {row['name'][:60]}", flush=True)

        if outcome is None:
            _write_row(ws, row["excel_row"], first_new_col, "—", "—", "—", "—", "Нет данных", "—")
            continue

        if outcome.local_hits:
            prices = [float(r.get("prise_per_unit")) for r, _ in outcome.local_hits
                      if r.get("prise_per_unit")]
        elif outcome.new_records:
            prices = [float(r.get("prise_per_unit")) for r in outcome.new_records
                      if r.get("prise_per_unit")]
        else:
            prices = [outcome.prediction.price] if outcome.prediction else []

        median = outcome.market_price
        min_p = min(prices) if prices else None
        max_p = max(prices) if prices else None
        source = ("локальная база" if outcome.local_hits else
                  "сайт zakup.sk.kz" if outcome.new_records else
                  "прогноз (RAG)" if outcome.prediction else "—")
        status = outcome.status or "Нет данных"
        confidence = (f"{outcome.prediction.confidence_pct:.0f}%"
                      if outcome.prediction else ("100%" if prices else "—"))

        _write_row(ws, row["excel_row"], first_new_col,
                   _fmt_price(median), _fmt_price(min_p), _fmt_price(max_p),
                   source, status, confidence)

    ws.freeze_panes = "A2"
    wb.save(output_path)
    log.info("Батч-режим: результаты сохранены в %s", output_path)
    print(f"Готово. Результаты сохранены в {output_path}")


def _write_row(ws, row_num: int, first_col: int, median: str, min_p: str, max_p: str,
               source: str, status: str, confidence: str) -> None:
    values = [median, min_p, max_p, source, status, confidence]
    for offset, value in enumerate(values):
        col = first_col + offset
        cell = ws.cell(row=row_num, column=col, value=value)
        cell.border = _BORDER
        cell.alignment = _CENTER
        if offset == 4:  # колонка "Статус"
            cell.fill = PatternFill("solid", fgColor=EXCEL_FILLS.get(status, EXCEL_FILLS["Нет данных"]))
            cell.font = _BOLD
